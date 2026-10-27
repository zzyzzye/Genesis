"""用真实 LangChain 客户端和隔离 HTTP 响应验证协议、文件与 checkpoint 回退。"""

from __future__ import annotations

import base64
import json
import re
from io import BytesIO
from typing import Any
from uuid import uuid4

import httpx
import httpx2
import pytest
from langchain_core.messages import HumanMessage
from langgraph.checkpoint.memory import InMemorySaver
from pydantic import SecretStr

from genesis_api.agent.runtime import EmbeddedAgentRuntime
from genesis_api.ai.schemas import AiChatRequest, AiConversationTitle, AiMessage
from genesis_api.billing.models import BillingOperation
from genesis_api.billing.transport import _validate_chat_budget
from genesis_api.core.config import Settings


def response_json(model: str) -> dict[str, Any]:
    return {
        "id": "resp-test", "object": "response", "created_at": 1, "status": "completed",
        "model": model, "output": [{
            "type": "message", "id": "msg-test", "role": "assistant", "status": "completed",
            "content": [{"type": "output_text", "text": "读取完成", "annotations": []}],
        }],
        "usage": {"input_tokens": 10, "output_tokens": 2, "total_tokens": 12},
    }


def chat_json(model: str) -> dict[str, Any]:
    return {
        "id": "chat-test", "object": "chat.completion", "created": 1, "model": model,
        "choices": [{"index": 0, "finish_reason": "stop", "message": {
            "role": "assistant", "content": "读取完成",
        }}],
        "usage": {"prompt_tokens": 10, "completion_tokens": 2, "total_tokens": 12},
    }


def docx_block() -> dict[str, Any]:
    from docx import Document

    document = Document()
    document.add_heading("回退参考文档", level=1)
    document.add_paragraph("通过 read_file 读取转换后的内容")
    buffer = BytesIO()
    document.save(buffer)
    return {
        "type": "file", "mime_type": (
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        ), "base64": base64.b64encode(buffer.getvalue()).decode(),
        "extras": {"filename": "参考.docx"},
    }


@pytest.mark.anyio
@pytest.mark.parametrize("provider", ["openai", "claude"])
@pytest.mark.parametrize("base_suffix", ["", "/v1"])
async def test_native_pdf_uses_provider_protocol(provider: str, base_suffix: str) -> None:
    """PDF 保留原始字节；Claude 使用 Messages，OpenAI 优先使用 Responses。"""
    sent: list[Any] = []
    model = "gpt-5.6-luna" if provider == "openai" else "claude-sonnet-4-6"

    def respond(request: Any) -> Any:
        sent.append(request)
        payload = response_json(model) if provider == "openai" else {
            "id": "msg-test", "type": "message", "role": "assistant", "model": model,
            "content": [{"type": "text", "text": "读取完成"}],
            "stop_reason": "end_turn", "stop_sequence": None,
            "usage": {"input_tokens": 10, "output_tokens": 2},
        }
        response_type = httpx.Response if provider == "openai" else httpx2.Response
        return response_type(200, json=payload)

    async with (
        httpx.AsyncClient(transport=httpx.MockTransport(respond)) as client,
        httpx2.AsyncClient(transport=httpx2.MockTransport(respond)) as anthropic_client,
    ):
        runtime = EmbeddedAgentRuntime()
        runtime._http_clients[True] = client
        runtime._anthropic_http_clients[True] = anthropic_client
        runtime._checkpointer = InMemorySaver()  # type: ignore[assignment]
        options: dict[str, Any] = {
            f"text_{provider}_api_key": SecretStr("synthetic"),
            f"text_{provider}_base_url": f"https://www.yyapi.cloud{base_suffix}",
        }
        settings = Settings(**options)
        graph = runtime._graph(settings, provider, model, "blog")
        await graph.ainvoke({"messages": [HumanMessage(content=[{
            "type": "file", "mime_type": "application/pdf", "base64": "JVBERi0=",
            "extras": {"filename": "参考.pdf"},
        }])]}, {"configurable": {"thread_id": str(uuid4())}})
    assert len(sent) == 1
    body = json.loads(sent[0].content)
    assert body["model"] == model
    assert sent[0].url.path == ("/v1/responses" if provider == "openai" else "/v1/messages")
    parts = body["input"] if provider == "openai" else body["messages"]
    expected_type = "input_file" if provider == "openai" else "document"
    block = next(block for message in parts if isinstance(message["content"], list)
                 for block in message["content"] if block.get("type") == expected_type)
    assert "JVBERi0=" in json.dumps(block)


@pytest.mark.anyio
@pytest.mark.parametrize("reject_protocol", [False, True])
async def test_rejected_protocol_or_file_uses_docling_and_persists_converted_message(
    reject_protocol: bool,
) -> None:
    """文件拒绝只转换文件；协议拒绝再换 Chat，转换文件能被原生工具读取并恢复。"""
    sent: list[httpx.Request] = []
    model = "gpt-5.6-luna"
    converted_path = ""

    def respond(request: httpx.Request) -> httpx.Response:
        nonlocal converted_path
        sent.append(request)
        body = json.loads(request.content)
        if len(sent) == 1:
            assert request.url.path == "/v1/responses"
            assert any(block.get("type") == "input_file" for message in body["input"]
                       if isinstance(message.get("content"), list)
                       for block in message["content"])
            return httpx.Response(
                404 if reject_protocol else 400,
                json={"detail": "Not Found"} if reject_protocol else {
                    "error": {"code": "unsupported_file_type", "message": "DOCX not supported"},
                },
            )
        assert request.url.path == ("/v1/chat/completions" if reject_protocol else "/v1/responses")
        serialized = json.dumps(body, ensure_ascii=False)
        assert "input_file" not in serialized and "file_data" not in serialized
        if len(sent) == 2:
            match = re.search(r"/attachments/[a-z0-9-]+/document.md", serialized)
            assert match is not None
            converted_path = match[0]
            if reject_protocol:
                payload = chat_json(model)
                payload["choices"][0].update({
                    "finish_reason": "tool_calls", "message": {"role": "assistant", "content": None,
                    "tool_calls": [{"id": "call-read", "type": "function", "function": {
                        "name": "read_file", "arguments": json.dumps({"file_path": converted_path}),
                    }}]},
                })
            else:
                payload = response_json(model)
                payload["output"] = [{
                    "type": "function_call", "id": "fc-read", "call_id": "call-read",
                    "name": "read_file", "arguments": json.dumps({"file_path": converted_path}),
                    "status": "completed",
                }]
        else:
            assert "读取转换后的内容" in serialized
            payload = chat_json(model) if reject_protocol else response_json(model)
        return httpx.Response(200, json=payload)

    async with httpx.AsyncClient(transport=httpx.MockTransport(respond)) as client:
        runtime = EmbeddedAgentRuntime()
        runtime._http_clients[True] = client
        runtime._checkpointer = InMemorySaver()  # type: ignore[assignment]
        graph = runtime._graph(Settings(
            text_openai_api_key=SecretStr("synthetic"),
            text_openai_base_url="https://www.yyapi.cloud",
        ), "openai", model, "blog")
        config = {"configurable": {"thread_id": str(uuid4())}}
        await graph.ainvoke({"messages": [HumanMessage(content=[docx_block()])]}, config)
        snapshot = await graph.aget_state(config)
        assert converted_path in snapshot.values["files"]
        assert all(
            block.get("type") != "file"
            for message in snapshot.values["messages"] if isinstance(message.content, list)
            for block in message.content if isinstance(block, dict)
        )
        await graph.ainvoke({"messages": [HumanMessage(content="继续")]}, config)
    assert len(sent) == 4


@pytest.mark.anyio
@pytest.mark.parametrize("status, error", [
    (401, {"error": {"code": "invalid_api_key"}}),
    (429, {"error": {"code": "rate_limit_exceeded"}}),
    (500, {"error": {"message": "upstream unavailable"}}),
    (404, {"error": {"code": "model_not_found", "message": "model not found"}}),
])
async def test_unrelated_errors_do_not_fallback(status: int, error: dict[str, Any]) -> None:
    from openai import APIStatusError

    sent: list[httpx.Request] = []

    def respond(request: httpx.Request) -> httpx.Response:
        sent.append(request)
        return httpx.Response(status, json=error)

    async with httpx.AsyncClient(transport=httpx.MockTransport(respond)) as client:
        runtime = EmbeddedAgentRuntime()
        runtime._http_clients[True] = client
        runtime._checkpointer = InMemorySaver()  # type: ignore[assignment]
        graph = runtime._graph(Settings(
            text_openai_api_key=SecretStr("synthetic"),
            text_openai_base_url="https://www.yyapi.cloud",
        ), "openai", "gpt-5.6-luna", "blog")
        with pytest.raises(APIStatusError):
            await graph.ainvoke({"messages": [HumanMessage(content="测试")]}, {
                "configurable": {"thread_id": str(uuid4())},
            })
    assert len(sent) == 1


@pytest.mark.anyio
@pytest.mark.parametrize("structured", [False, True])
async def test_structured_generation_and_title_share_protocol_fallback(structured: bool) -> None:
    """无 Agent 的结构化生成及标题也能在端点拒绝后回退，且保持原生契约校验。"""
    sent: list[httpx.Request] = []
    model = "gpt-5.6-luna"

    def respond(request: httpx.Request) -> httpx.Response:
        sent.append(request)
        if request.url.path == "/v1/responses":
            return httpx.Response(404, json={"detail": "Not Found"})
        assert request.url.path == "/v1/chat/completions"
        body = json.loads(request.content)
        content = json.dumps({"title": "协议回退"}, ensure_ascii=False)
        if body.get("stream"):
            # 契约完整前分片不能被提前校验或作为最终结果返回。
            split = content.index("回退")
            chunks = [{
                "id": "chat-title", "object": "chat.completion.chunk", "created": 1,
                "model": model, "choices": [{"index": 0, "finish_reason": None,
                "delta": {"role": "assistant", "content": content[:split]}}],
            }, {
                "id": "chat-title", "object": "chat.completion.chunk", "created": 1,
                "model": model, "choices": [{"index": 0, "finish_reason": None,
                "delta": {"content": content[split:]}}],
            }, {
                "id": "chat-title", "object": "chat.completion.chunk", "created": 1,
                "model": model, "choices": [{"index": 0, "finish_reason": "stop", "delta": {}}],
            }]
            return httpx.Response(200, headers={"content-type": "text/event-stream"}, content=(
                "".join(f"data: {json.dumps(chunk)}\n\n" for chunk in chunks)
                + "data: [DONE]\n\n"
            ))
        payload = chat_json(model)
        payload["choices"][0]["message"]["content"] = content
        return httpx.Response(200, json=payload)

    async with httpx.AsyncClient(transport=httpx.MockTransport(respond)) as client:
        runtime = EmbeddedAgentRuntime()
        runtime._http_clients[True] = client
        settings = Settings(
            text_openai_api_key=SecretStr("synthetic"), text_openai_model=model,
            text_openai_base_url="https://www.yyapi.cloud",
        )
        if structured:
            result = await runtime.generate_structured(
                settings, prompt="概括主题", content="测试", schema=AiConversationTitle,
            )
            assert result.title == "协议回退"
        else:
            assert await runtime.generate_title(AiChatRequest(
                surface="studio", messages=[AiMessage(role="user", content="测试")],
            ), settings) == "协议回退"
    assert len(sent) == 2


@pytest.mark.parametrize("builtin", [False, True])
def test_paid_anthropic_request_accepts_only_custom_tools(builtin: bool) -> None:
    """切换原生 Messages 后仍核验预算，且不会放行有额外费用的服务端工具。"""
    settings = Settings(
        text_claude_api_key=SecretStr("synthetic"),
        text_claude_base_url="https://www.yyapi.cloud",
    )
    operation = BillingOperation(
        provider="claude", model="claude-sonnet-4-6", policy={"max_output_tokens": 4096},
    )
    tool: dict[str, Any] = {"name": "read_file", "input_schema": {"type": "object"}}
    if builtin:
        tool["type"] = "web_search_20250305"
    request = httpx2.Request("POST", "https://www.yyapi.cloud/v1/messages", json={
        "model": operation.model, "max_tokens": 4096, "tools": [tool],
    })
    if builtin:
        with pytest.raises(RuntimeError, match="供应商内置工具"):
            _validate_chat_budget(operation, settings, request)
    else:
        _validate_chat_budget(operation, settings, request)
