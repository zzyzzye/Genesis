"""附件权限、真实文档解析、原生文件导入及标准多模态输入验收。"""

from __future__ import annotations

import base64
from collections.abc import Iterator
from datetime import UTC, datetime
from io import BytesIO
from pathlib import Path
from typing import Any
from uuid import UUID, uuid4

import pytest
from fastapi import HTTPException
from langchain_core.messages import ToolMessage
from langchain_core.runnables import RunnableConfig
from sqlalchemy import create_engine
from sqlalchemy.orm import Session

from genesis_api.agent.documents import document_text
from genesis_api.agent.files import AttachmentFilesMiddleware, agent_files_context
from genesis_api.ai.conversations import conversation_detail
from genesis_api.ai.models import AiChatRun, AiConversation
from genesis_api.ai.runs import create_ai_chat_run
from genesis_api.ai.schemas import AiAttachment, AiChatRequest, AiMessage
from genesis_api.blog import attachments as storage
from genesis_api.blog.agent.attachments import attachment_input, prepare_attachments
from genesis_api.core.config import Settings
from genesis_api.database.base import Base
from genesis_api.identity.models import User, UserRole
from genesis_api.llm.capabilities import ModelCapabilities


@pytest.fixture
def session(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Iterator[Session]:
    """仅访问测试内存数据库和临时附件目录。"""
    monkeypatch.setattr(storage, "STORAGE_ROOT", tmp_path)
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    with Session(engine, expire_on_commit=False) as value:
        yield value
    engine.dispose()


def owner(session: Session, name: str = "owner") -> User:
    user = User(handle=name, display_name=name, role=UserRole.OWNER)
    session.add(user)
    session.commit()
    return user


def file(session: Session, user_id: UUID, name: str, content: bytes) -> UUID:
    from genesis_api.blog.models import BlogAttachment

    item_id = uuid4()
    session.add(
        BlogAttachment(
            id=item_id,
            owner_id=user_id,
            name=name,
            kind="file",
            size=len(content),
            created_at=datetime.now(UTC),
            updated_at=datetime.now(UTC),
        )
    )
    session.commit()
    (storage.STORAGE_ROOT / str(item_id)).write_bytes(content)
    return item_id


def request(user_id: UUID, ids: list[UUID], **kwargs: Any) -> AiChatRequest:
    return AiChatRequest(
        surface="studio",
        actor_id=user_id,
        actor_role="owner",
        provider="openai",
        model="test",
        messages=[AiMessage(role="user", content="", attachment_ids=ids)],
        **kwargs,
    )


def test_owned_document_metadata_is_rebuilt_and_persisted(session: Session) -> None:
    user = owner(session)
    item_id = file(session, user.id, "参考.md", "# 真实资料\n用于写作".encode())
    payload = request(user.id, [item_id, item_id])
    payload.messages[0].attachments = [
        AiAttachment(
            id=uuid4(),
            name="客户端伪造.pdf",
            size=123,
            mime_type="application/pdf",
            kind="document",
        )
    ]
    created = create_ai_chat_run(session, user_id=user.id, request=payload, settings=Settings())
    run = session.get(AiChatRun, created.id)
    assert run is not None
    assert payload.messages[0].attachment_ids == [item_id]
    assert payload.messages[0].attachments[0].name == "参考.md"
    assert (
        AiChatRequest.model_validate(run.request_payload).messages[0].attachments[0].name
        == "参考.md"
    )
    conversation = AiConversation(user_id=user.id)
    session.add(conversation)
    session.flush()
    run.conversation_id = conversation.id
    session.commit()
    history = conversation_detail(session, conversation)
    assert history.messages[0].attachments[0].id == item_id
    messages, files = attachment_input(session, user.id, payload)
    assert files[0][1].decode() == "# 真实资料\n用于写作"
    assert files[0][0] in str(messages)
    assert "用于写作" not in str(messages)


def test_foreign_or_missing_attachment_is_not_exposed(session: Session) -> None:
    user, other = owner(session), owner(session, "other")
    for item_id in [file(session, other.id, "私有.txt", b"private"), uuid4()]:
        with pytest.raises(HTTPException) as caught:
            prepare_attachments(session, user.id, request(user.id, [item_id]), Settings())
        assert caught.value.status_code == 404


@pytest.mark.parametrize("name", ["参考.png", "参考", "参考.jpg"])
def test_chat_media_keeps_filename_in_native_block_and_prompt(session: Session, name: str) -> None:
    """聊天按真实内容补后缀；模型适配器能接收标准块且保留文件身份。"""
    from langchain_core.messages import HumanMessage
    from langchain_google_genai.chat_models import _convert_to_parts
    from langchain_openai.chat_models.base import _convert_message_to_dict

    user = owner(session)
    content = b"\x89PNG\r\n\x1a\n" + b"0" * 40
    item_id = file(session, user.id, name, content)
    messages, files = attachment_input(session, user.id, request(user.id, [item_id]))
    assert not files
    blocks = messages[0]["content"]
    assert isinstance(blocks, list)
    assert "参考.png" in blocks[0]["text"]
    assert blocks[1]["extras"]["filename"] == "参考.png"
    assert blocks[1]["mime_type"] == "image/png"
    assert base64.b64decode(blocks[1]["base64"]) == content
    converted = _convert_message_to_dict(HumanMessage(content=blocks))
    assert converted["content"][1]["image_url"]["url"].startswith("data:image/png;base64,")
    parts = _convert_to_parts(blocks)
    assert parts[1].inline_data is not None
    assert parts[1].inline_data.mime_type == "image/png"


@pytest.mark.parametrize(
    "name, content",
    [
        ("empty.txt", b""),
        ("payload.svg", b"<svg/>"),
        ("archive.zip", b"PK\x03\x04"),
    ],
)
def test_empty_or_unsupported_files_are_rejected(
    session: Session,
    name: str,
    content: bytes,
) -> None:
    user = owner(session)
    item_id = file(session, user.id, name, content)
    with pytest.raises(HTTPException):
        prepare_attachments(session, user.id, request(user.id, [item_id]), Settings())


def test_model_capabilities_and_video_transport_are_enforced(
    session: Session,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from genesis_api.blog.agent import attachments

    user = owner(session)
    image = file(session, user.id, "image.png", b"\x89PNG\r\n\x1a\n" + b"0" * 40)
    video = file(session, user.id, "video.mp4", b"\x00\x00\x00\x18ftypmp42" + b"0" * 40)
    monkeypatch.setattr(
        attachments, "model_capabilities_for", lambda *_: ModelCapabilities(profile={})
    )
    with pytest.raises(HTTPException) as caught:
        prepare_attachments(session, user.id, request(user.id, [image]), Settings())
    assert "图片" in caught.value.detail
    monkeypatch.setattr(
        attachments,
        "model_capabilities_for",
        lambda *_: ModelCapabilities(
            profile={"image_inputs": True, "video_inputs": True},
        ),
    )
    with pytest.raises(HTTPException):
        prepare_attachments(session, user.id, request(user.id, [video]), Settings())
    payload = request(user.id, [image, video])
    payload.provider = "gemini"
    payload = prepare_attachments(
        session,
        user.id,
        payload,
        Settings(
            text_gemini_base_url="https://generativelanguage.googleapis.com",
        ),
    )
    messages, files = attachment_input(session, user.id, payload)
    assert not files
    blocks = messages[0]["content"]
    assert isinstance(blocks, list)
    assert blocks[1]["type"] == "image"
    assert base64.b64decode(blocks[2]["base64"]).startswith(b"\x00\x00\x00\x18ftyp")
    # 验证现有供应商适配器实际接受标准块，不仅检查我们生成的字典。
    from langchain_core.messages import HumanMessage
    from langchain_google_genai.chat_models import _convert_to_parts
    from langchain_openai.chat_models.base import _convert_message_to_dict

    parts = _convert_to_parts(blocks)
    assert parts[2].inline_data is not None
    assert parts[2].inline_data.mime_type == "video/mp4"
    converted = _convert_message_to_dict(HumanMessage(content=[blocks[1]]))
    assert converted["content"][0]["type"] == "image_url"


def test_media_attachments_reuse_private_loader(
    session: Session, monkeypatch: pytest.MonkeyPatch,
) -> None:
    """影音使用同一鉴权附件入口，不复制解析或公开文件。"""
    from genesis_api.agent.capabilities import agent_capabilities
    from genesis_api.blog.agent import capability

    user = owner(session)
    item_id = file(session, user.id, "分镜.txt", "雨夜街道".encode())
    payload = request(user.id, [item_id])
    payload.surface = "media"
    prepared = prepare_attachments(session, user.id, payload, Settings())
    assert prepared.messages[0].attachments[0].id == item_id
    monkeypatch.setattr(capability, "SessionLocal", lambda: session)
    inputs, files = agent_capabilities.load_attachments("media", prepared)
    assert inputs[0]["role"] == "user"
    assert files[0][1].decode() == "雨夜街道"
    prepared.actor_role = "member"
    with pytest.raises(RuntimeError, match="权限"):
        agent_capabilities.load_attachments("media", prepared)


def test_total_size_role_scope_and_changed_file_are_rejected(session: Session) -> None:
    from genesis_api.blog.agent import attachments
    from genesis_api.blog.models import BlogAttachment

    user = owner(session)
    item_id = file(session, user.id, "文本.txt", b"test")
    payload = request(user.id, [item_id])
    payload.surface = "toolbox"
    with pytest.raises(HTTPException):
        prepare_attachments(session, user.id, payload, Settings())
    payload.surface = "studio"
    payload.actor_role = "member"
    with pytest.raises(HTTPException):
        prepare_attachments(session, user.id, payload, Settings())
    payload.actor_role = "owner"
    record = session.get(BlogAttachment, item_id)
    assert record is not None
    record.size = attachments.MAX_BYTES + 1
    with pytest.raises(HTTPException):
        prepare_attachments(session, user.id, payload, Settings())
    record.size = 4
    (storage.STORAGE_ROOT / str(item_id)).write_bytes(b"changed")
    with pytest.raises(RuntimeError, match="无法读取附件"):
        attachment_input(session, user.id, payload)


def test_docx_and_pdf_use_real_loaders(session: Session) -> None:
    """真实 DOCX 标题、表格和 PDF 正文经 Docling 读取，兼容 UUID 存储路径。"""
    from docx import Document
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    user = owner(session)
    docx = BytesIO()
    document = Document()
    document.add_heading("参考文档", level=1)
    document.add_paragraph("DOCX attachment")
    table = document.add_table(rows=2, cols=2)
    for row, values in zip(table.rows, [("人物", "场景"), ("苏晚", "雨夜")], strict=True):
        for cell, value in zip(row.cells, values, strict=True):
            cell.text = value
    document.save(docx)
    pdf = BytesIO()
    writer = PdfWriter()
    page = writer.add_blank_page(width=400, height=400)
    page[NameObject("/Resources")] = DictionaryObject(
        {
            NameObject("/Font"): DictionaryObject(
                {
                    NameObject("/F1"): DictionaryObject(
                        {
                            NameObject("/Type"): NameObject("/Font"),
                            NameObject("/Subtype"): NameObject("/Type1"),
                            NameObject("/BaseFont"): NameObject("/Helvetica"),
                        }
                    ),
                }
            )
        }
    )
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 40 300 Td (PDF attachment) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(stream)
    writer.write(pdf)
    ids = [
        file(session, user.id, "文档.docx", docx.getvalue()),
        file(session, user.id, "文档.pdf", pdf.getvalue()),
    ]
    messages, files = attachment_input(session, user.id, request(user.id, ids))
    assert not files
    assert len(messages[0]["content"]) == 3  # type: ignore[arg-type]
    markdown = document_text(
        storage.STORAGE_ROOT / str(ids[0]),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    assert "DOCX attachment" in markdown
    assert "# 参考文档" in markdown
    assert "| 人物" in markdown
    assert "苏晚" in markdown and "雨夜" in markdown
    assert "PDF attachment" in document_text(storage.STORAGE_ROOT / str(ids[1]), "application/pdf")


@pytest.mark.parametrize("encrypted", [False, True])
def test_docling_rejects_excessive_pages_and_encrypted_pdf(
    session: Session, encrypted: bool,
) -> None:
    """页数上限在转换前执行，加密文档不绕过读取限制。"""
    from docling.exceptions import ConversionError
    from pypdf import PdfWriter

    writer = PdfWriter()
    for _ in range(1 if encrypted else 201):
        writer.add_blank_page(width=100, height=100)
    if encrypted:
        writer.encrypt("test-only")
    pdf = BytesIO()
    writer.write(pdf)
    user = owner(session)
    item_id = file(session, user.id, "受限.pdf", pdf.getvalue())
    with pytest.raises((ValueError, ConversionError)):
        document_text(storage.STORAGE_ROOT / str(item_id), "application/pdf")


def test_historical_media_prevents_switch_to_text_only_model(
    session: Session,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from genesis_api.blog.agent import attachments

    user = owner(session)
    conversation = AiConversation(user_id=user.id)
    session.add(conversation)
    session.flush()
    item_id = file(session, user.id, "image.png", b"\x89PNG\r\n\x1a\n" + b"0" * 40)
    payload = request(user.id, [item_id], conversation_id=conversation.id)
    monkeypatch.setattr(
        attachments,
        "model_capabilities_for",
        lambda *_: ModelCapabilities(
            profile={"image_inputs": True},
        ),
    )
    create_ai_chat_run(session, user_id=user.id, request=payload, settings=Settings())
    next_turn = payload.model_copy(update={"messages": [AiMessage(role="user", content="继续")]})
    monkeypatch.setattr(
        attachments, "model_capabilities_for", lambda *_: ModelCapabilities(profile={})
    )
    with pytest.raises(HTTPException) as error:
        prepare_attachments(session, user.id, next_turn, Settings())
    assert "当前对话含图片" in error.value.detail


def test_count_and_total_size_limits_are_per_request(session: Session) -> None:
    from genesis_api.blog.models import BlogAttachment

    user = owner(session)
    ids = [file(session, user.id, f"资料-{index}.txt", b"test") for index in range(9)]
    payload = request(user.id, ids[:8])
    payload.messages.append(AiMessage(role="user", attachment_ids=ids[8:]))
    with pytest.raises(HTTPException, match="8 个附件"):
        prepare_attachments(session, user.id, payload, Settings())
    for item_id in ids[:2]:
        record = session.get(BlogAttachment, item_id)
        assert record is not None
        record.size = 12 * 1024 * 1024
    with pytest.raises(HTTPException, match="20 MiB"):
        prepare_attachments(session, user.id, request(user.id, ids[:2]), Settings())


def test_invalid_document_and_empty_pdf_have_actionable_error(session: Session) -> None:
    from pypdf import PdfWriter

    user = owner(session)
    pdf = BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.write(pdf)
    item_id = file(session, user.id, "broken.docx", b"PKbroken")
    with pytest.raises(RuntimeError, match="无法读取附件"):
        attachment_input(session, user.id, request(user.id, [item_id]))
    blank = file(session, user.id, "blank.pdf", pdf.getvalue())
    with pytest.raises(ValueError, match="未提取到文字"):
        document_text(storage.STORAGE_ROOT / str(blank), "application/pdf")
    invalid = file(session, user.id, "encoding.txt", b"ok\n\xff")
    with pytest.raises(HTTPException):
        prepare_attachments(session, user.id, request(user.id, [invalid]), Settings())


@pytest.mark.anyio
async def test_native_upload_files_is_thread_scoped_and_checkpointed() -> None:
    from deepagents import create_deep_agent
    from langchain_core.language_models.fake_chat_models import FakeMessagesListChatModel
    from langchain_core.messages import AIMessage
    from langgraph.checkpoint.memory import InMemorySaver

    class Model(FakeMessagesListChatModel):
        def bind_tools(self, *args: Any, **kwargs: Any) -> Model:
            return self

    path = "/attachments/test/reference.txt"
    model = Model(
        responses=[
            AIMessage(
                content="",
                tool_calls=[
                    {
                        "name": "read_file",
                        "args": {"file_path": path},
                        "id": "read-1",
                        "type": "tool_call",
                    }
                ],
            ),
            AIMessage(content="读取完成"),
        ]
    )
    graph = create_deep_agent(
        model=model, checkpointer=InMemorySaver(), middleware=[AttachmentFilesMiddleware()]
    )
    config: RunnableConfig = {"configurable": {"thread_id": "one"}}
    with agent_files_context([(path, "原生文件导入验收".encode())]):
        result = await graph.ainvoke(
            {"messages": [{"role": "user", "content": "读取参考文档"}]}, config
        )
    assert any(
        isinstance(msg, ToolMessage) and "原生文件导入验收" in str(msg.content)
        for msg in result["messages"]
    )
    assert path in result["files"]
    again = await graph.ainvoke({"messages": [{"role": "user", "content": "再次读取"}]}, config)
    assert path in again["files"]
    other = await graph.ainvoke(
        {"messages": [{"role": "user", "content": "其他会话"}]},
        {"configurable": {"thread_id": "two"}},
    )
    assert path not in other.get("files", {})
